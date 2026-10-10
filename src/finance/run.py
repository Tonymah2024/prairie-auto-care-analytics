"""Pipeline step 'finance': ledger -> statements -> ratios -> forecast & scenarios -> tables/DB -> Excel (verified by an Excel
recalculation when Microsoft Excel is available) -> validation. Everything is SIMULATED and labelled."""
from __future__ import annotations

import json
import shutil
import sqlite3
import tempfile
from dataclasses import replace
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

from src.config import DB_PATH, FIN_OUT, OUT, TABLES
from src.finance import excel as X
from src.finance import forecast as F
from src.finance import ledger as LG
from src.finance import policy as P
from src.finance import ratios as R
from src.finance import statements as S
from src.finance import validation as V

MODEL = FIN_OUT / "Prairie_Auto_Care_Integrated_Financial_Model.xlsx"


def _csv(df, name, status="SIMULATED", index=True):
    d = df.copy()
    if "data_status" not in d.columns:
        d["data_status"] = status
    d.to_csv(TABLES / f"fin_{name}.csv", index=index)


def insurer_revenue_monthly(T):
    inv = T["customer_invoices"]
    x = inv[inv.bill_to_type == "Insurer"]
    return x.groupby(x.invoice_date.dt.to_period("M").astype(str)).subtotal_cad.sum()


def compute():
    L = LG.build()
    St = S.build(L)
    rat = R.compute(St["is_m"], St["bs_m"], St["bs_open"], insurer_revenue_monthly(L["T"]))
    ann = R.annual_table(rat)
    FC = F.build(L, St)
    return L, St, rat, ann, FC


def tables(L, St, rat, ann, FC, Vr):
    for per, tag in (("m", "monthly"), ("q", "quarterly"), ("y", "annual")):
        _csv(St[f"is_{per}"].rename_axis("key").assign(label=lambda d: d.index.map(S.IS_LABEL)), f"income_statement_{tag}")
        _csv(St[f"bs_{per}"].rename_axis("key").assign(label=lambda d: d.index.map(S.BS_LABEL)), f"balance_sheet_{tag}")
        _csv(St[f"cf_{per}"].rename_axis("key").assign(label=lambda d: d.index.map(S.CF_LABEL)), f"cash_flow_{tag}")
    _csv(St["cs_m"].rename_axis("key").assign(label=lambda d: d.index.map(S.IS_LABEL)), "common_size_monthly")
    _csv(St["cs_y"].rename_axis("key").assign(label=lambda d: d.index.map(S.IS_LABEL)), "common_size_annual")
    _csv(St["bs_open"].to_frame("2023-12").rename_axis("key"), "opening_balance_sheet")
    _csv(St["direct"], "cash_flow_direct_method")
    _csv(St["tb"], "trial_balance_monthly", index=False)
    _csv(pd.DataFrame(P.COA, columns=["code", "name", "type", "normal_balance", "statement_line", "cash_flow_class"]), "chart_of_accounts", index=False)
    _csv(St["subledger"], "subledger_reconciliation", index=False)
    _csv(St["far"], "fixed_asset_register", index=False)
    _csv(L["treasury"], "debt_treasury_schedule", index=False)
    _csv(St["inv_val"], "inventory_valuation")
    _csv(St["wip_rec"], "wip_reconciliation", index=False)
    _csv(St["branch_is"], "branch_income_statement_monthly", index=False)
    _csv(St["branch_summary"].reset_index(), "branch_pnl_annual", index=False)
    _csv(St["branch_bal"], "branch_working_capital_monthly", index=False)
    _csv(St["branch_cf"], "branch_operating_cash_flow_monthly", index=False)
    _csv(pd.DataFrame([(k, *v) for k, v in P.ALLOCATION.items()], columns=["account", "driver", "explanation"]), "allocation_assumptions", index=False)
    _csv(R.ar_aging_monthly(L["ar_items"], LG.MONTHS), "ar_aging_monthly", index=False)
    for nm, items in (("ar", L["ar_items"]), ("ap", L["ap_items"])):
        ag = S.aging(items)
        _csv(ag.assign(doc_date=ag.doc_date.astype(str), due_date=ag.due_date.astype(str)).drop(columns=[c for c in ("closed_date",) if c in ag]), f"{nm}_open_items_30sep2026", index=False)
    _csv(rat.rename_axis("ratio"), "ratios_monthly_ttm")
    _csv(ann.rename_axis("key"), "ratios_annual")
    for n, r in FC["results"].items():
        _csv(r.rename_axis("key").assign(label=lambda d: d.index.map(lambda k: F.ROW[k][1])), f"forecast_{n.lower()}_monthly", "FORECAST (SIMULATED)")
    _csv(FC["summary"], "forecast_scenario_summary", "FORECAST (SIMULATED)")
    _csv(FC["sensitivity"], "forecast_sensitivity", "FORECAST (SIMULATED)", index=False)
    _csv(FC["base"]["drivers"].rename_axis("driver"), "forecast_drivers", "FORECAST INPUT (SIMULATED)")
    _csv(pd.Series(FC["base"]["A"]).rename("value").to_frame().rename_axis("assumption"), "forecast_assumptions", "FORECAST INPUT (SIMULATED)")
    _csv(Vr["checks"], "validation_checks", index=False)
    _csv(Vr["revenue_cost"], "revenue_cost_reconciliation", index=False)
    _csv(Vr["branch_rec"], "branch_consolidation_reconciliation", index=False)
    js = {"A": FC["base"]["A"], "drivers": {k: list(map(float, v)) for k, v in FC["base"]["drivers"].iterrows()},
          "opening": {k: float(v) for k, v in FC["base"]["opening"].items()}, "data_status": "FORECAST INPUT (SIMULATED)"}
    (TABLES / "fin_forecast_base.json").write_text(json.dumps(js, indent=1))
    gl = L["gl"].copy()
    gl["date"], gl["month"] = gl.date.dt.strftime("%Y-%m-%d"), gl.month.astype(str)
    gl.to_csv(FIN_OUT / "general_ledger.csv.gz", index=False, compression="gzip")
    je = gl.groupby("je_id").agg(date=("date", "first"), source=("source", "first"), reference=("reference", "first"), description=("description", "first"),
                                 lines=("line_no", "count"), total_debits=("debit", "sum"), total_credits=("credit", "sum")).reset_index().assign(data_status="SIMULATED")
    je.to_csv(FIN_OUT / "journal_entries.csv.gz", index=False, compression="gzip")
    con = sqlite3.connect(DB_PATH)
    for name, df in [("fin_chart_of_accounts", pd.DataFrame(P.COA, columns=["code", "name", "type", "normal_balance", "statement_line", "cash_flow_class"])),
                     ("fin_journal_entries", je), ("fin_general_ledger", gl), ("fin_trial_balance", St["tb"]),
                     ("fin_ar_subledger", L["ar_items"].astype({"doc_date": str, "due_date": str, "closed_date": str})), ("fin_ap_subledger", L["ap_items"].astype({"doc_date": str, "due_date": str})),
                     ("fin_branch_income_statement", St["branch_is"]), ("fin_wip_reconciliation", St["wip_rec"]),
                     ("fin_fixed_assets", St["far"]), ("fin_debt_treasury", L["treasury"])]:
        df.assign(data_status="SIMULATED").to_sql(name, con, if_exists="replace", index=False)
    con.close()


def verify_excel(model, base):
    if not X.excel_available():
        return {"ran": False, "ok": None}
    tmp = Path(tempfile.mkdtemp(prefix="pac_xl_"))
    keys = ["net_rev", "gross_profit", "net_income", "cash", "wip", "ar_insurer", "coll_hours", "total_assets", "total_le", "cfo", "tax"]
    cust = {"eff_change": 0.10, "appt_change": 0.12, "parts_cost_change": -0.05, "cycle_change": -0.15, "collect_change": -0.2, "hires": 2.0, "opex_change": 0.04, "rate_change": 0.02}
    cases = [("Base", None), ("Conservative", None), ("Custom", cust)]
    rows, worst, all_ok = [], 0.0, True
    for name, edits_ in cases:
        dst = tmp / f"{name}.xlsx"
        edits, b, sc = {}, base, F.SCENARIOS.get(name, F.SCENARIOS["Base"])
        if edits_:
            for fld, v in edits_.items():
                edits[("Assumptions", "G" + model.apos[fld].split("$")[-1])] = v
            edits[("Drivers", "H" + str(model.pos["capex_d"][1]))] = 60000.0              # Feb-2027 capex
            dv = base["drivers"].copy()
            dv.loc["capex_d", 5] = 60000.0
            b, sc = {"A": base["A"], "drivers": dv, "opening": base["opening"]}, replace(F.SCENARIOS["Base"], **edits_)
        X.set_scenario(MODEL, dst, name, edits)
        if not X.recalc_with_excel(dst, tmp):
            return {"ran": False, "ok": None, "error": "Excel recalculation failed"}
        wb = openpyxl.load_workbook(dst, data_only=True)
        ws = wb["Forecast"]
        py = F.run(sc, b)
        diff = max(float(np.abs(np.array([ws.cell(row=model.pos[k][1], column=3 + t).value or 0.0 for t in range(1, F.H + 1)], float) - py.loc[k].iloc[1:].values).max()) for k in keys)
        ok = bool(wb["Checks"][model.all_checks.split("!")[1]].value)
        worst, all_ok = max(worst, diff), all_ok and ok
        rows.append({"case": name, "rows_compared": len(keys) * F.H, "max_abs_diff_vs_python": diff, "excel_checks_all_true": ok,
                     "excel_net_income_12m": float(sum(ws.cell(row=model.pos["net_income"][1], column=3 + t).value for t in range(1, F.H + 1)))})
        if name == "Base":
            shutil.copy(dst, MODEL)
    shutil.rmtree(tmp, ignore_errors=True)
    return {"ran": True, "ok": worst < 0.01 and all_ok, "max_diff": worst, "excel_checks": all_ok, "cases": "Base; Conservative; Custom", "rows": rows}


def run():
    L, St, rat, ann, FC = compute()
    Vr = V.run(L, St, FC)
    tables(L, St, rat, ann, FC, Vr)
    model = X.Model(L, St, FC)
    model.build(MODEL)
    X.statements_workbook(L, St, FIN_OUT / "Prairie_Auto_Care_Financial_Statements.xlsx")
    X.ratios_workbook(rat, ann, FIN_OUT / "Prairie_Auto_Care_Financial_Ratios.xlsx")
    X.working_capital_workbook(St, rat, FC, FIN_OUT / "Prairie_Auto_Care_Working_Capital_Analysis.xlsx")
    X.forecast_workbook(FC, FIN_OUT / "Prairie_Auto_Care_12M_Forecast_Scenarios.xlsx")
    xl = verify_excel(model, FC["base"])
    (TABLES / "fin_excel_verification.json").write_text(json.dumps({**xl, "data_status": "SIMULATED"}, indent=1, default=str))
    Vr = V.run(L, St, FC, excel_results=xl, out_dir=OUT)
    _csv(Vr["checks"], "validation_checks", index=False)
    failed = Vr["checks"][Vr["checks"].result != "PASS"]
    if len(failed):
        raise SystemExit(f"Financial validation failed:\n{failed.to_string()}")
    return {"checks_passed": int((Vr["checks"].result == "PASS").sum()), "excel": {k: v for k, v in xl.items() if k != "rows"}}


if __name__ == "__main__":
    print(run())
